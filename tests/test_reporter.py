from datetime import datetime

from openpyxl import load_workbook

from consultor.models import NO_VERIFICADO, POSIBLE_NOVEDAD, SIN_CAMBIO
from consultor.reporter import LEYENDA, avisar, escribir_reporte, leer_decisiones

AHORA = datetime(2026, 10, 12, 7, 0, 0)
RESUMEN = {"total": 3, "exitosas": 2, "fallidas": 1, "novedades": 1, "estado": "Completo"}


def fila(**kw):
    base = dict(
        alerta_id=None,
        empresa="ICEIN",
        radicado="11001400307720210114700",
        despacho="JUZGADO 077",
        estado_consulta="Exitosa",
        resultado=SIN_CAMBIO,
        anterior="",
        fecha_detectada="",
        detectada="",
        anotacion="",
        hora="2026-10-12T07:00:00",
        motivo="",
        decision="",
        validada_por="",
        validada_en="",
    )
    base.update(kw)
    return base


def novedad(alerta_id=7, **kw):
    return fila(
        alerta_id=alerta_id,
        resultado=POSIBLE_NOVEDAD,
        detectada="Auto",
        fecha_detectada="2026-10-10",
        decision="Pendiente",
        **kw,
    )


def escribir(tmp_path, filas):
    return escribir_reporte(filas, RESUMEN, tmp_path, 3, AHORA)


def valores(hoja):
    return [[c.value for c in f] for f in hoja.iter_rows()]


def test_reporte_tiene_tres_hojas_y_la_leyenda(tmp_path):
    ruta = escribir(tmp_path, [fila()])
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Alertas", "Sin cambio"]
    textos = [str(c.value) for f in wb["Resumen"].iter_rows() for c in f]
    assert LEYENDA in textos
    assert ruta.name == "reporte_ciclo_0003_2026-10-12.xlsx"


def test_novedades_primero_y_sin_cambio_en_su_hoja(tmp_path):
    filas = [fila(), fila(resultado=NO_VERIFICADO, motivo="sin resultados"), novedad()]
    wb = load_workbook(escribir(tmp_path, filas))
    enc = valores(wb["Alertas"])[0]
    resultados = [f[enc.index("Resultado")] for f in valores(wb["Alertas"])[1:]]
    assert resultados == [POSIBLE_NOVEDAD, NO_VERIFICADO]
    sin = valores(wb["Sin cambio"])
    assert [f[enc.index("Resultado")] for f in sin[1:]] == [SIN_CAMBIO]


def test_formula_en_la_anotacion_se_guarda_como_texto(tmp_path):
    wb = load_workbook(escribir(tmp_path, [novedad(anotacion="=1+1")]))
    hoja = wb["Alertas"]
    enc = [c.value for c in hoja[1]]
    celda = hoja.cell(row=2, column=enc.index("Anotación detectada") + 1)
    assert celda.value == "=1+1"
    assert celda.data_type == "s"


def test_no_incluye_nombres_de_partes(tmp_path):
    wb = load_workbook(escribir(tmp_path, [novedad()]))
    encabezados = [c.value for c in wb["Alertas"][1]]
    assert not any("parte" in str(e).lower() or "sujeto" in str(e).lower() for e in encabezados)


def marcar(ruta, decision, valido_por):
    wb = load_workbook(ruta)
    hoja = wb["Alertas"]
    enc = [c.value for c in hoja[1]]
    hoja.cell(row=2, column=enc.index("Decisión") + 1, value=decision)
    hoja.cell(row=2, column=enc.index("Validó") + 1, value=valido_por)
    wb.save(ruta)


def test_leer_decisiones_devuelve_solo_las_decididas(tmp_path):
    ruta = escribir(tmp_path, [novedad(alerta_id=7)])
    assert leer_decisiones(tmp_path) == []  # todo Pendiente
    marcar(ruta, "Confirmada", "Alisson Rengifo")
    assert leer_decisiones(tmp_path) == [(7, "Confirmada", "Alisson Rengifo")]


def test_validador_en_blanco_se_devuelve_vacio(tmp_path):
    ruta = escribir(tmp_path, [novedad(alerta_id=7)])
    marcar(ruta, "Descartada", None)
    assert leer_decisiones(tmp_path) == [(7, "Descartada", "")]


def test_archivo_ilegible_avisa_y_sigue(tmp_path):
    (tmp_path / "reporte_ciclo_0001_2026-10-10.xlsx").write_bytes(b"no es un excel")
    ruta = escribir(tmp_path, [novedad(alerta_id=7)])
    marcar(ruta, "Confirmada", "Alisson Rengifo")
    avisos = []
    assert leer_decisiones(tmp_path, avisos.append) == [(7, "Confirmada", "Alisson Rengifo")]
    assert any("reporte_ciclo_0001" in a for a in avisos)


def test_avisar_escribe_en_el_log(tmp_path, capsys):
    archivo = tmp_path / "carpeta" / "avisos.log"
    avisar("Ciclo 3 listo", archivo, AHORA)
    assert "2026-10-12 07:00:00 Ciclo 3 listo" in archivo.read_text(encoding="utf-8")
    assert "Ciclo 3 listo" in capsys.readouterr().out

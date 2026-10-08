from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook

import consultor.main as main_mod
from consultor.bloqueo import Bloqueo
from consultor.main import MOTIVO_DETENIDO, correr_ciclo
from consultor.models import ERROR, FALLIDA, Consulta, Radicado
from consultor.reporter import LEYENDA
from consultor.store import Store
from tests.utiles import Reloj, act, ok

R1 = Radicado("11001400307720210114700", empresa="ICEIN")


class FetcherFalso:
    def __init__(self, respuestas):
        self.respuestas = respuestas
        self.llamadas = []

    def consultar(self, r):
        self.llamadas.append(r.radicado)
        resp = self.respuestas[r.radicado]
        if isinstance(resp, Exception):
            raise resp
        return resp

    def pausar(self):
        pass


def correr(tmp_path, radicados, fetcher, store, reloj, avisos=None, max_fallas=0.5, progreso=None):
    avisos = [] if avisos is None else avisos
    return correr_ciclo(
        radicados, fetcher, store, tmp_path, "Alisson Rengifo", max_fallas, avisos.append, reloj, progreso
    )


def hoja_alertas(ruta):
    hoja = load_workbook(ruta)["Alertas"]
    enc = [c.value for c in hoja[1]]
    filas = [[c.value for c in f] for f in hoja.iter_rows(min_row=2)]
    return enc, filas


def dos_ciclos(tmp_path):
    store, reloj = Store(":memory:"), Reloj()
    f = FetcherFalso({R1.radicado: ok(act(1), act(2))})
    r1 = correr(tmp_path, [R1], f, store, reloj)
    reloj.avanzar()
    f.respuestas = {R1.radicado: ok(act(3, texto="Auto"), act(2), act(1))}
    r2 = correr(tmp_path, [R1], f, store, reloj)
    return store, reloj, f, r1, r2


def test_primera_vez_sin_alerta_y_luego_detecta_la_nueva(tmp_path):
    avisos = []
    store, reloj = Store(":memory:"), Reloj()
    f = FetcherFalso({R1.radicado: ok(act(1), act(2))})
    r1 = correr(tmp_path, [R1], f, store, reloj, avisos)
    assert (r1.total, r1.novedades, r1.estado) == (1, 0, "Completo")

    reloj.avanzar()
    f.respuestas = {R1.radicado: ok(act(3, texto="Auto"), act(2), act(1))}
    r2 = correr(tmp_path, [R1], f, store, reloj, avisos)
    assert r2.novedades == 1
    enc, filas = hoja_alertas(r2.reporte)
    assert filas[0][enc.index("Resultado")] == "POSIBLE NOVEDAD"
    assert filas[0][enc.index("Actuación detectada")] == "Auto"
    assert LEYENDA in avisos[-1]


def test_decision_de_alisson_se_registra_y_no_se_repite(tmp_path):
    store, reloj, f, r1, r2 = dos_ciclos(tmp_path)
    wb = load_workbook(r2.reporte)
    hoja = wb["Alertas"]
    enc = [c.value for c in hoja[1]]
    hoja.cell(row=2, column=enc.index("Decisión") + 1, value="Descartada")
    wb.save(r2.reporte)

    reloj.avanzar()
    r3 = correr(tmp_path, [R1], f, store, reloj)
    assert r3.novedades == 0  # R6: misma actuación, no alerta de nuevo
    anterior = store.filas_reporte(r2.ciclo_id)[0]
    assert anterior["decision"] == "Descartada"
    assert anterior["validada_por"] == "Alisson Rengifo"  # validador por defecto (CA5)


def test_r5_detiene_y_marca_el_resto_como_no_verificado(tmp_path):
    radicados = [Radicado(f"{i:023d}") for i in range(1, 13)]
    caida = Consulta(FALLIDA, motivo="HTTP 503", falla_portal=True)
    f = FetcherFalso({r.radicado: caida for r in radicados})
    store, avisos = Store(":memory:"), []
    res = correr(tmp_path, radicados, f, store, Reloj(), avisos)
    assert res.estado.startswith("Detenido")
    assert len(f.llamadas) == 10
    assert (res.total, res.exitosas, res.fallidas) == (12, 0, 12)
    assert any("Fuente no disponible" in a for a in avisos)
    motivos = [x["motivo"] for x in store.filas_reporte(res.ciclo_id)]
    assert motivos.count(MOTIVO_DETENIDO) == 2


def test_ciclo_interrumpido_se_reanuda_sin_repetir_consultas(tmp_path):
    rs = [Radicado(f"{i:023d}") for i in range(1, 4)]
    store, reloj = Store(":memory:"), Reloj()
    f = FetcherFalso(
        {rs[0].radicado: ok(act(1)), rs[1].radicado: ok(act(1)), rs[2].radicado: RuntimeError("red caída")}
    )
    with pytest.raises(RuntimeError):
        correr(tmp_path, rs, f, store, reloj)
    f2 = FetcherFalso({rs[2].radicado: ok(act(1))})
    res = correr(tmp_path, rs, f2, store, reloj)
    assert f2.llamadas == [rs[2].radicado]
    assert res.total == 3


def test_radicado_sin_resultados_queda_en_el_reporte_como_no_verificado(tmp_path):
    f = FetcherFalso({R1.radicado: Consulta(FALLIDA, motivo="sin resultados")})
    res = correr(tmp_path, [R1], f, Store(":memory:"), Reloj())
    enc, filas = hoja_alertas(res.reporte)
    assert filas[0][enc.index("Resultado")] == "NO VERIFICADO"
    assert filas[0][enc.index("Motivo")] == "sin resultados"
    assert (res.total, res.exitosas, res.fallidas) == (1, 0, 1)


def test_si_falla_el_reporte_el_ciclo_queda_abierto_y_se_reanuda_sin_reconsultar(tmp_path, monkeypatch):
    store, reloj = Store(":memory:"), Reloj()
    f = FetcherFalso({R1.radicado: ok(act(1))})
    correr(tmp_path, [R1], f, store, reloj)
    reloj.avanzar()
    f.respuestas = {R1.radicado: ok(act(2, texto="Auto"), act(1))}
    original = main_mod.escribir_reporte

    def falla(*a, **k):
        raise RuntimeError("disco lleno")

    monkeypatch.setattr(main_mod, "escribir_reporte", falla)
    with pytest.raises(RuntimeError):
        correr(tmp_path, [R1], f, store, reloj)
    monkeypatch.setattr(main_mod, "escribir_reporte", original)
    llamadas_antes = len(f.llamadas)
    res = correr(tmp_path, [R1], f, store, reloj)
    assert len(f.llamadas) == llamadas_antes  # no reconsulta
    assert res.novedades == 1
    enc, filas = hoja_alertas(res.reporte)
    assert filas[0][enc.index("Resultado")] == "POSIBLE NOVEDAD"


def test_ciclo_interrumpido_de_otro_dia_se_avisa_y_su_alerta_sigue_visible(tmp_path, monkeypatch):
    store, reloj, avisos = Store(":memory:"), Reloj(), []
    f = FetcherFalso({R1.radicado: ok(act(1))})
    correr(tmp_path, [R1], f, store, reloj, avisos)
    reloj.avanzar()
    f.respuestas = {R1.radicado: ok(act(2, texto="Auto"), act(1))}

    def falla(*a, **k):
        raise RuntimeError("se cayó la máquina")

    monkeypatch.undo()
    monkeypatch.setattr(main_mod, "escribir_reporte", falla)
    with pytest.raises(RuntimeError):
        correr(tmp_path, [R1], f, store, reloj, avisos)
    monkeypatch.undo()
    reloj.avanzar()
    res = correr(tmp_path, [R1], f, store, reloj, avisos)
    assert any("interrumpido" in a.lower() for a in avisos)
    enc, filas = hoja_alertas(res.reporte)
    assert [x[enc.index("Actuación detectada")] for x in filas] == ["Auto"]


def test_portal_caido_con_pocos_radicados_no_cierra_como_completo(tmp_path):
    radicados = [Radicado(f"{i:023d}") for i in range(1, 4)]
    caida = Consulta(FALLIDA, motivo="HTTP 503", falla_portal=True)
    f = FetcherFalso({r.radicado: caida for r in radicados})
    avisos = []
    res = correr(tmp_path, radicados, f, Store(":memory:"), Reloj(), avisos)
    assert res.estado == "Completo con fallas de la fuente"
    assert any("Fuente no disponible" in a for a in avisos)


FORMA_ROTA = Consulta(ERROR, motivo="respuesta inesperada: KeyError", falla_portal=True)


def test_un_error_de_forma_cierra_con_alerta_de_cambio_en_la_api(tmp_path):
    rs = [Radicado(f"{i:023d}") for i in range(1, 6)]
    resp = {r.radicado: ok(act(1)) for r in rs}
    resp[rs[4].radicado] = FORMA_ROTA
    avisos = []
    res = correr(tmp_path, rs, FetcherFalso(resp), Store(":memory:"), Reloj(), avisos)
    assert res.estado == "Completo con alerta: posible cambio en la API"
    assert any("posiblemente cambió" in a and "1 de 5" in a for a in avisos)
    assert not any("Fuente no disponible" in a for a in avisos)


def test_muchos_errores_de_forma_detienen_con_alerta_de_api(tmp_path):
    rs = [Radicado(f"{i:023d}") for i in range(1, 13)]
    avisos = []
    res = correr(tmp_path, rs, FetcherFalso({r.radicado: FORMA_ROTA for r in rs}), Store(":memory:"), Reloj(), avisos)
    assert res.estado == "Detenido: posible cambio en la API"
    assert any("posiblemente cambió" in a for a in avisos)


def test_todo_correcto_no_lanza_alerta_de_api(tmp_path):
    rs = [Radicado(f"{i:023d}") for i in range(1, 4)]
    avisos = []
    res = correr(tmp_path, rs, FetcherFalso({r.radicado: ok(act(1)) for r in rs}), Store(":memory:"), Reloj(), avisos)
    assert res.estado == "Completo"
    assert not any("posiblemente cambió" in a for a in avisos)


def test_lista_vacia_avisa_y_no_es_completo(tmp_path):
    avisos = []
    res = correr(tmp_path, [], FetcherFalso({}), Store(":memory:"), Reloj(), avisos)
    assert res.estado == "Lista vacía"
    assert res.reporte is None
    assert any("vacía" in a for a in avisos)


def test_el_progreso_se_informa_por_cada_radicado(tmp_path):
    rs = [Radicado(f"{i:023d}") for i in range(1, 4)]
    f = FetcherFalso({r.radicado: ok(act(1)) for r in rs})
    visto = []
    correr(tmp_path, rs, f, Store(":memory:"), Reloj(), progreso=lambda hecho, total, rad: visto.append((hecho, total, rad)))
    assert visto == [(1, 3, rs[0].radicado), (2, 3, rs[1].radicado), (3, 3, rs[2].radicado)]


def preparar(tmp_path, monkeypatch, minutos=30):
    wb = Workbook()
    ws = wb.active
    ws.title = "GENERAL"
    ws.append(["Radicado"])
    ws.append([R1.radicado])
    wb.save(tmp_path / "x.xlsx")
    cfg = tmp_path / "config.toml"
    cfg.write_text(
        f'[fuente]\ntipo = "excel"\nruta = "{(tmp_path / "x.xlsx").as_posix()}"\nhoja = "GENERAL"\ncol_radicado = "Radicado"\n'
        f'[salida]\nbase_datos = "{(tmp_path / "datos" / "c.db").as_posix()}"\n'
        f'carpeta_reportes = "{(tmp_path / "reportes").as_posix()}"\nvalidador = "Alisson"\n'
        '[portal]\npausa_segundos = 0\nreintentos = 1\nespera_segundos = 0\ntimeout_segundos = 5\nmax_fallas_ciclo = 0.5\n'
        f'[ejecucion]\nmin_minutos_entre_ciclos = {minutos}\n',
        encoding="utf-8",
    )

    class Falso:
        def __init__(self, **k):
            pass

        def consultar(self, r):
            return ok(act(1))

        def pausar(self):
            pass

    monkeypatch.setattr(main_mod, "Fetcher", Falso)
    return ["run", "--config", str(cfg)]


def test_un_segundo_intento_seguido_es_rechazado_con_mensaje_claro(tmp_path, monkeypatch, capsys):
    args = preparar(tmp_path, monkeypatch)
    assert main_mod.ejecutar(args) == 0
    assert main_mod.ejecutar(args) == 3
    assert "espera hasta las" in capsys.readouterr().out


def test_forzar_salta_el_limite(tmp_path, monkeypatch):
    args = preparar(tmp_path, monkeypatch)
    assert main_mod.ejecutar(args) == 0
    assert main_mod.ejecutar(args + ["--forzar"]) == 0


def test_con_limite_en_cero_no_rechaza(tmp_path, monkeypatch):
    args = preparar(tmp_path, monkeypatch, minutos=0)
    assert main_mod.ejecutar(args) == 0
    assert main_mod.ejecutar(args) == 0


def test_si_hay_otro_ciclo_corriendo_se_rechaza_y_dice_su_avance(tmp_path, monkeypatch, capsys):
    args = preparar(tmp_path, monkeypatch)
    with Bloqueo(tmp_path / "datos") as b:
        b.publicar("06:30 | 12 de 44")
        assert main_mod.ejecutar(args) == 3
    assert "12 de 44" in capsys.readouterr().out


def test_solo_radicado_no_cuenta_para_el_limite(tmp_path, monkeypatch):
    args = preparar(tmp_path, monkeypatch)
    assert main_mod.ejecutar(args + ["--solo-radicado", R1.radicado]) == 0
    assert main_mod.ejecutar(args) == 0  # no fue rechazado


ESTADO_PAUSADO = "Pausado: el portal bloqueó las consultas"
BLOQUEO = Consulta(FALLIDA, motivo="HTTP 403: el portal bloqueó las consultas", falla_portal=True, bloqueo=True)


def cinco():
    return [Radicado(f"{i:023d}") for i in range(1, 6)]


def con_bloqueo_en_el_tercero(rs):
    resp = {r.radicado: ok(act(1)) for r in rs}
    resp[rs[2].radicado] = BLOQUEO
    resp[rs[3].radicado] = RuntimeError("no debe consultarse")
    resp[rs[4].radicado] = RuntimeError("no debe consultarse")
    return FetcherFalso(resp)


def test_un_bloqueo_detiene_el_ciclo_avisa_y_deja_pendientes(tmp_path):
    rs, avisos = cinco(), []
    f = con_bloqueo_en_el_tercero(rs)
    res = correr(tmp_path, rs, f, Store(":memory:"), Reloj(), avisos)
    assert len(f.llamadas) == 3
    assert res.estado == ESTADO_PAUSADO
    assert res.pendientes == 3
    assert res.total == 2
    alerta = [a for a in avisos if "bloqueó" in a]
    assert len(alerta) == 1 and "Quedan 3" in alerta[0]
    assert alerta[0] == (
        "ALERTA: el portal bloqueó las consultas (HTTP 403 o 429) después de 2 radicados. "
        "Quedan 3 sin consultar. La próxima ejecución continuará desde ahí; espere al menos 30 minutos."
    )
    assert not any("Fuente no disponible" in a for a in avisos)


def test_el_minimo_de_minutos_se_incluye_en_el_aviso(tmp_path):
    rs, avisos = cinco(), []
    correr_ciclo(
        rs, con_bloqueo_en_el_tercero(rs), Store(":memory:"), tmp_path, "A", 0.5, avisos.append, Reloj(), min_minutos=45
    )
    assert any("al menos 45 minutos" in a for a in avisos)


def test_tras_un_bloqueo_se_reanuda_solo_con_los_pendientes(tmp_path):
    rs, store, reloj = cinco(), Store(":memory:"), Reloj()
    r1 = correr(tmp_path, rs, con_bloqueo_en_el_tercero(rs), store, reloj)
    assert store.ultimo_cierre() is not None  # el límite entre ejecuciones actúa de enfriamiento
    reloj.actual += timedelta(minutes=35)
    f2 = FetcherFalso({r.radicado: ok(act(1)) for r in rs[2:]})
    r2 = correr(tmp_path, rs, f2, store, reloj)
    assert f2.llamadas == [r.radicado for r in rs[2:]]
    assert r2.ciclo_id == r1.ciclo_id
    assert r2.estado == "Completo" and r2.total == 5 and r2.pendientes == 0


def test_el_reporte_del_bloqueo_trae_los_pendientes(tmp_path):
    rs = cinco()
    res = correr(tmp_path, rs, con_bloqueo_en_el_tercero(rs), Store(":memory:"), Reloj())
    assert res.reporte.exists()
    resumen = {f[0]: f[1] for f in load_workbook(res.reporte)["Resumen"].iter_rows(values_only=True)}
    assert resumen["Pendientes por consultar"] == 3
    assert resumen["Estado del ciclo"] == ESTADO_PAUSADO


def test_un_ciclo_normal_no_muestra_pendientes_en_el_reporte(tmp_path):
    res = correr(tmp_path, [R1], FetcherFalso({R1.radicado: ok(act(1))}), Store(":memory:"), Reloj())
    resumen = {f[0] for f in load_workbook(res.reporte)["Resumen"].iter_rows(values_only=True)}
    assert "Pendientes por consultar" not in resumen

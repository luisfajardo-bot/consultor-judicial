import re

from openpyxl import load_workbook

from .models import Radicado


def _cortar(fuente: str) -> list[str]:
    """Radicados de 23 dígitos (o varios pegados) en los tramos de dígitos del texto."""
    return [
        tramo[i : i + 23]
        for tramo in re.findall(r"[0-9]+", fuente)
        if len(tramo) >= 23 and len(tramo) % 23 == 0
        for i in range(0, len(tramo), 23)
    ]


def _radicados(valor) -> list[str]:
    """Radicados de una celda. El cuadro trae texto junto al número ("SEGUNDAINSTANCIA...")
    y a veces varios radicados pegados o separados. Se hace en tres pasos, sin perder ninguno:

    1) En el texto tal cual, cortando cada tramo de dígitos de a 23: "…700-2026" no pega el año.
    2) En lo que el paso 1 no usó, sin espacios, guiones ni puntos: "11001-40-03-077-2021-01147-00".
    3) Un tramo largo que no es un radicado completo (un dígito de más o de menos) se devuelve
       tal cual: se ve como NO VERIFICADO y no se pierde en silencio. Años, fechas, juzgados,
       NIT y teléfonos tienen menos de 18 dígitos.
    Si no hay nada de eso se devuelve el texto limpio, para que se vea como NO VERIFICADO.
    """
    if valor is None:
        return []
    if isinstance(valor, float):  # Excel guarda solo 15 cifras: el número ya está recortado
        return [f"numérico: {valor!r}"]
    texto = str(valor)
    encontrados = _cortar(texto)
    resto = re.sub(
        r"[0-9]+",
        lambda m: " " if len(m.group()) >= 23 and len(m.group()) % 23 == 0 else m.group(),
        texto,
    )
    limpio_resto = re.sub(r"[\s\-\.]", "", resto)
    encontrados += _cortar(limpio_resto)
    encontrados += [t for t in re.findall(r"[0-9]+", limpio_resto) if len(t) >= 18 and len(t) % 23 != 0]
    limpio = re.sub(r"[\s\-\.]", "", texto)
    return encontrados or ([limpio] if limpio else [])


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _indice(encabezado: list[str], nombre: str, hoja: str) -> int:
    try:
        return encabezado.index(nombre)
    except ValueError:
        raise ValueError(f"No encuentro la columna '{nombre}' en la hoja '{hoja}'") from None


class FuenteExcel:
    """Lee los radicados de un Excel, solo lectura (R9).

    Contrato de cualquier fuente: un objeto con cargar() -> list[Radicado].
    No descarta radicados inválidos: el fetcher los marca NO VERIFICADO (R3).
    """

    def __init__(
        self,
        ruta,
        hoja,
        col_radicado,
        col_empresa="",
        col_despacho="",
        col_calidad="",
        col_estado="",
        estados_incluidos=(),
    ):
        self.ruta = ruta
        self.hoja = hoja
        self.col_radicado = col_radicado
        self.col_empresa = col_empresa
        self.col_despacho = col_despacho
        self.col_calidad = col_calidad
        self.col_estado = col_estado
        self.estados = tuple(e.strip().lower() for e in estados_incluidos if e.strip())
        self.filas_sin_radicado = 0

    def cargar(self) -> list[Radicado]:
        self.filas_sin_radicado = 0
        wb = load_workbook(self.ruta, read_only=True, data_only=True)
        try:
            filas = wb[self.hoja].iter_rows(values_only=True)
            enc = [_texto(c) for c in next(filas)]
            i_rad = _indice(enc, self.col_radicado, self.hoja)
            i_emp = _indice(enc, self.col_empresa, self.hoja) if self.col_empresa else None
            i_des = _indice(enc, self.col_despacho, self.hoja) if self.col_despacho else None
            i_cal = _indice(enc, self.col_calidad, self.hoja) if self.col_calidad else None
            i_est = _indice(enc, self.col_estado, self.hoja) if self.estados else None

            def celda(fila, i):
                return _texto(fila[i]) if i is not None and i < len(fila) else ""

            por_radicado: dict[str, Radicado] = {}
            for fila in filas:
                # "ACTIVO" admite variantes del Excel como "ACTIVO -COBRO COSTAS"
                if self.estados and not celda(fila, i_est).lower().startswith(self.estados):
                    continue
                rads = _radicados(fila[i_rad]) if i_rad < len(fila) else []
                if not rads:
                    if any(_texto(c) for j, c in enumerate(fila) if j != i_est):
                        self.filas_sin_radicado += 1  # datos sin radicado: se avisa, no se pierde en silencio
                    continue
                for rad in rads:
                    nuevo = Radicado(rad, celda(fila, i_emp), celda(fila, i_des), celda(fila, i_cal))
                    previo = por_radicado.get(rad)
                    if previo is None:
                        por_radicado[rad] = nuevo
                    elif nuevo.empresa and nuevo.empresa not in previo.empresa.split(" / "):
                        empresas = " / ".join(e for e in (previo.empresa, nuevo.empresa) if e)
                        por_radicado[rad] = Radicado(rad, empresas, previo.despacho, previo.calidad)
            return list(por_radicado.values())
        finally:
            wb.close()


def crear_fuente(cfg: dict):
    tipo = cfg.get("tipo")
    if tipo == "excel":
        return FuenteExcel(
            ruta=cfg["ruta"],
            hoja=cfg["hoja"],
            col_radicado=cfg["col_radicado"],
            col_empresa=cfg.get("col_empresa", ""),
            col_despacho=cfg.get("col_despacho", ""),
            col_calidad=cfg.get("col_calidad", ""),
            col_estado=cfg.get("col_estado", ""),
            estados_incluidos=cfg.get("estados_incluidos", ()),
        )
    raise ValueError(f"fuente desconocida: {tipo}")

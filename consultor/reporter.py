import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .models import CONFIRMADA, DESCARTADA, NO_VERIFICADO, POSIBLE_NOVEDAD, SIN_CAMBIO

FUENTE = "Consulta de Procesos Nacional Unificada"
ENLACE_PORTAL = "https://consultaprocesos.ramajudicial.gov.co/Procesos/NumeroRadicacion"
LEYENDA = (
    "Alerta automática. No constituye notificación procesal ni actuación confirmada; "
    "verificar en la fuente oficial."
)

ENCABEZADOS = [
    "ID alerta",
    "Empresa",
    "Radicado",
    "Despacho",
    "Estado de la consulta",
    "Resultado",
    "Actuación anterior",
    "Fecha actuación detectada",
    "Actuación detectada",
    "Anotación detectada",
    "Fecha y hora de consulta",
    "Fuente",
    "Enlace",
    "Motivo",
    "Decisión",
    "Validó",
    "Validada en",
]
ORDEN = {POSIBLE_NOVEDAD: 0, NO_VERIFICADO: 1, SIN_CAMBIO: 2}


def _valores(f: dict) -> list:
    return [
        f["alerta_id"],
        f["empresa"],
        f["radicado"],
        f["despacho"],
        f["estado_consulta"],
        f["resultado"],
        f["anterior"],
        f["fecha_detectada"],
        f["detectada"],
        f["anotacion"],
        f["hora"],
        FUENTE,
        ENLACE_PORTAL,
        f["motivo"],
        f["decision"],
        f["validada_por"],
        f["validada_en"],
    ]


def _agregar(hoja, valores: list) -> None:
    hoja.append(valores)
    for celda in hoja[hoja.max_row]:
        # un texto que empieza por "=" se guardaría como fórmula
        if isinstance(celda.value, str) and celda.value.startswith("="):
            celda.data_type = "s"


def escribir_reporte(filas, resumen: dict, carpeta, ciclo_id: int, ahora: datetime) -> Path:
    wb = Workbook()
    hoja_resumen = wb.active
    hoja_resumen.title = "Resumen"
    for linea in [
        ("Ciclo", ciclo_id),
        ("Fecha", ahora.strftime("%Y-%m-%d %H:%M")),
        ("Total consultados", resumen["total"]),
        ("Exitosos", resumen["exitosas"]),
        ("Fallidos", resumen["fallidas"]),
        ("Posibles novedades", resumen["novedades"]),
        ("Estado del ciclo", resumen["estado"]),
        ("", ""),
        ("Aviso", LEYENDA),
    ]:
        hoja_resumen.append(linea)

    alertas = wb.create_sheet("Alertas")
    sin_cambio = wb.create_sheet("Sin cambio")
    alertas.append(ENCABEZADOS)
    sin_cambio.append(ENCABEZADOS)
    for f in sorted(filas, key=lambda f: ORDEN[f["resultado"]]):
        _agregar(sin_cambio if f["resultado"] == SIN_CAMBIO else alertas, _valores(f))

    col = ENCABEZADOS.index("Decisión") + 1
    letra = alertas.cell(row=1, column=col).column_letter
    lista = DataValidation(type="list", formula1='"Pendiente,Confirmada,Descartada"', allow_blank=True)
    alertas.add_data_validation(lista)
    lista.add(f"{letra}2:{letra}{max(alertas.max_row, 500)}")
    alertas.freeze_panes = "A2"
    sin_cambio.freeze_panes = "A2"

    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"reporte_ciclo_{ciclo_id:04d}_{ahora:%Y-%m-%d}.xlsx"
    wb.save(ruta)
    return ruta


def leer_decisiones(carpeta, avisar_fn=None) -> list[tuple[int, str, str]]:
    """Lee las decisiones que Alisson marcó en los reportes anteriores.

    Devuelve (alerta_id, decisión, validó). Un archivo ilegible (por ejemplo abierto
    en Excel) se avisa y se reintenta en el ciclo siguiente.
    """
    decisiones = []
    for ruta in sorted(Path(carpeta).glob("reporte_*.xlsx")):
        try:
            wb = load_workbook(ruta, read_only=True, data_only=True)
            try:
                filas = wb["Alertas"].iter_rows(values_only=True)
                enc = next(filas)
                i_id = enc.index("ID alerta")
                i_dec = enc.index("Decisión")
                i_val = enc.index("Validó")
                for f in filas:
                    if f[i_id] and f[i_dec] in (CONFIRMADA, DESCARTADA):
                        decisiones.append((int(f[i_id]), f[i_dec], f[i_val] or ""))
            finally:
                wb.close()
        except (OSError, ValueError, KeyError, StopIteration, zipfile.BadZipFile) as e:
            if avisar_fn:
                avisar_fn(
                    f"No se pudo leer {ruta.name} ({type(e).__name__}); "
                    "se reintenta en el próximo ciclo."
                )
    return decisiones


def avisar(texto: str, archivo, ahora: datetime) -> None:
    """Único punto de aviso. El canal real (correo o Teams) lo define Sistemas."""
    archivo = Path(archivo)
    archivo.parent.mkdir(parents=True, exist_ok=True)
    with open(archivo, "a", encoding="utf-8") as f:
        f.write(f"{ahora:%Y-%m-%d %H:%M:%S} {texto}\n")
    print(texto)
